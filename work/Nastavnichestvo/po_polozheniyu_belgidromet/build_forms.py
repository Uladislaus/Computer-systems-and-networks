#!/usr/bin/env python3
"""Документы наставничества по Положению Белгидромета (приложения 1–3 из PDF кадров)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "docx"
OUT.mkdir(parents=True, exist_ok=True)

ORG = "Белгидромет"
ORG_FULL = (
    "государственном учреждении «Республиканский центр по гидрометеорологии, "
    "контролю радиоактивного загрязнения и мониторингу окружающей среды»"
)
UNIT = "Служба программного обеспечения"
WORKER = "Руцкий Владислав Анатольевич"
WORKER_POS = "инженер-программист"
MENTOR = "Чаусов Олег Викторович"
MENTOR_POS = "ведущий инженер-программист"
BOSS = "Войтешонок Александр Леонидович"
BOSS_POS = "начальник отдела разработки программного обеспечения"
PERIOD = "с 18.08.2026 по 18.10.2026"
BASIS = "приказ № 154-ОД от 14.05.2026; Положение о наставничестве Белгидромета"


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def para(doc, text="", *, size=12, bold=False, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, italic=italic)
    return p


def new_doc():
    d = Document()
    s = d.sections[0]
    s.top_margin = Cm(1.5)
    s.bottom_margin = Cm(1.5)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(1.5)
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return d


def signatures(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    set_font(p.add_run("__________          ________________          ________________"), 12)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(8)
    set_font(p2.add_run("дата                         подпись                              (Ф.И.О.)"), 9, italic=True)


def build_agreement() -> Path:
    doc = new_doc()
    para(doc, "Приложение 1 к Положению о наставничестве в\n" + ORG_FULL, size=10, align=WD_ALIGN_PARAGRAPH.RIGHT, after=12)
    para(doc, "СОГЛАШЕНИЕ", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(doc, "о сотрудничестве", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(doc, "между наставником и работником", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)

    para(doc, f"Я, наставник, {MENTOR_POS} {MENTOR},")
    para(doc, f"беру на обучение {WORKER_POS}а {WORKER}")
    para(
        doc,
        "и обязуюсь помочь овладеть профессией, повысить образовательный и культурный уровень, "
        "способствовать скорейшей адаптации в коллективе.",
        after=8,
    )
    signatures(doc)
    para(doc, f"                                                                 {MENTOR}", size=10, after=14)

    para(doc, f"Я, работник, {WORKER_POS} {WORKER},")
    para(
        doc,
        "обязуюсь овладеть необходимыми профессиональными навыками, добросовестно выполнять "
        "поставленные цели и задачи, максимально использовать полученные от наставника опыт и знания "
        "для повышения профессионального уровня.",
        after=8,
    )
    signatures(doc)
    para(doc, f"                                                                 {WORKER}", size=10, after=10)
    para(doc, f"Основание: {BASIS}.", size=10, italic=True, after=2)
    para(doc, f"Период наставничества: {PERIOD}.", size=10, italic=True)

    path = OUT / "01_Soglashenie_o_sotrudnichestve.docx"
    doc.save(path)
    return path


def build_program() -> Path:
    doc = new_doc()
    para(doc, "Приложение 2 к Положению о наставничестве в\n" + ORG_FULL, size=10, align=WD_ALIGN_PARAGRAPH.RIGHT, after=8)
    para(doc, "УТВЕРЖДАЮ", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(doc, f"_________________ / {BOSS} /", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(doc, f"({BOSS_POS})", size=9, italic=True, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(doc, "«____» ______________ 202__ г.", align=WD_ALIGN_PARAGRAPH.RIGHT, after=12)

    para(doc, "ПРОГРАММА НАСТАВНИЧЕСТВА", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(doc, f"в структурном подразделении: {UNIT}", align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(doc, f"({ORG})", align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    rows = [
        ("1", "Отбор / закрепление наставника", "август 2026", ""),
        ("2", "Подготовка наставника, методическая поддержка", "по мере необходимости", ""),
        ("3", "Определение сроков наставничества", PERIOD, ""),
        ("4", f"Закрепление наставника ({MENTOR}) за работником ({WORKER})", "не позднее 5 раб. дней", ""),
        ("5", "Утверждение индивидуальных планов наставничества", "в течение 10 раб. дней", ""),
        ("6", "Обмен опытом / консультации наставника", "еженедельно", ""),
        ("7", "Информационное сопровождение (кадры, локальные акты)", "в течение периода", ""),
        ("8", "Промежуточный контроль выполнения плана", "сентябрь 2026", ""),
        ("9", "Итоговое собеседование, отчёт и характеристика", "октябрь 2026", ""),
    ]
    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = "Table Grid"
    headers = ["№ п/п", "Основные мероприятия*", "Дата", "Ответственные осуществления наставничества (дата, подпись)"]
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(h), size=9, bold=True)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(val), size=9)
    para(doc, "", after=6)
    para(
        doc,
        "* Перечень мероприятий конкретизируется с учётом особенностей проведения наставничества "
        "в организации (структурном подразделении).",
        size=9,
        italic=True,
    )
    path = OUT / "02_Programma_nastavnichestva.docx"
    doc.save(path)
    return path


def build_plan() -> Path:
    doc = new_doc()
    para(doc, "Приложение 3 к Положению о наставничестве в\n" + ORG_FULL, size=10, align=WD_ALIGN_PARAGRAPH.RIGHT, after=8)
    para(doc, "УТВЕРЖДАЮ", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(doc, f"_________________ / {BOSS} /", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    para(
        doc,
        "(руководитель структурного подразделения либо лицо,\nответственное за организацию проведения наставничества)",
        size=9,
        italic=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
        after=0,
    )
    para(doc, "«____» ______________ 202__ г.", align=WD_ALIGN_PARAGRAPH.RIGHT, after=12)

    para(doc, "ИНДИВИДУАЛЬНЫЙ ПЛАН НАСТАВНИЧЕСТВА", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    para(doc, f"Ф.И.О. работника: {WORKER}", after=2)
    para(doc, f"профессия (должность): {WORKER_POS}", after=2)
    para(doc, f"структурное подразделение: {UNIT}", after=2)
    para(doc, f"Наставник: {MENTOR_POS} {MENTOR}", after=2)
    para(doc, f"Период наставничества: {PERIOD}", after=10)

    events = [
        ("1", "Составление и утверждение индивидуального плана наставничества", "в течение 10 рабочих дней", ""),
        ("2", "Ознакомление с основными целями и задачами, направлениями текущей деятельности Белгидромета и службы ПО", "до 27.08.2026", ""),
        ("3", "Ознакомление со структурой организации, её подразделениями, установленными взаимосвязями и подчинённостью", "до 27.08.2026", ""),
        ("4", "Детальное изучение деятельности структурного подразделения (структура, подчинённость, полномочия, регламент работы)", "до 03.09.2026", ""),
        ("5", "Изучение нормативных документов по направлению деятельности (локальные акты, информационная безопасность, доступы к ИС)", "до 03.09.2026", ""),
        ("6", "Освоение профессиональных навыков: рабочее окружение Python (venv, pip, Git, IDE)", "август 2026", ""),
        ("7", "Освоение профессиональных навыков: язык Python (скрипты, отладка, стиль кода) на рабочих примерах наставника", "август–сентябрь 2026", ""),
        ("8", "Освоение профессиональных навыков: веб-фреймворк Django (проект, ORM, migrate, CRUD, базовые настройки безопасности)", "сентябрь 2026", ""),
        ("9", "Освоение профессиональных навыков: СУБД PostgreSQL (запросы, схема БД, связка с Django)", "сентябрь 2026", ""),
        ("10", "Освоение профессиональных навыков: СУБД Oracle (подключение, типовые запросы, отличия от PostgreSQL)", "сентябрь–октябрь 2026", ""),
        ("11", "Освоение профессиональных навыков: вводный уровень нейронных сетей / ML на Python применительно к задачам подразделения", "октябрь 2026", ""),
        ("12", "Интеграция: участие в текущих задачах службы ПО под контролем наставника; подготовка материалов к итоговому отчёту", "октябрь 2026", ""),
        ("13", "Участие в собеседовании с руководителем структурного подразделения по итогам завершения наставничества", "до 18.10.2026", ""),
        ("14", "Подведение итогов выполнения индивидуального плана наставничества", "до 18.10.2026", ""),
    ]

    table = doc.add_table(rows=1 + len(events), cols=4)
    table.style = "Table Grid"
    for j, h in enumerate(["№ п/п", "Мероприятия¹", "Срок исполнения", "Отчёт о выполнении² (дата, подпись)"]):
        cell = table.rows[0].cells[j]
        cell.text = ""
        set_font(cell.paragraphs[0].add_run(h), size=9, bold=True)
    for i, row in enumerate(events, start=1):
        for j, val in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            set_font(cell.paragraphs[0].add_run(val), size=9)

    para(doc, "", after=4)
    para(
        doc,
        "¹ Перечень мероприятий конкретизируется с учётом особенностей проведения наставничества "
        "в организации (структурном подразделении).",
        size=9,
        italic=True,
        after=2,
    )
    para(
        doc,
        "² Заполняется наставником в свободной форме (например: выполнено, не выполнено, требуется повторное изучение и т.д.).",
        size=9,
        italic=True,
        after=12,
    )

    signatures(doc)
    para(doc, f"                                                                 {MENTOR} (наставник)", size=10, after=4)
    signatures(doc)
    para(doc, f"                                                                 {WORKER} (работник)", size=10, after=12)

    para(doc, "Оценка выполнения индивидуального плана наставничества: ________________", after=8)
    para(doc, "Характеристика на работника:", bold=True, after=4)
    for _ in range(4):
        para(doc, "_" * 85, after=2)
    para(doc, "", after=4)
    signatures(doc)
    para(doc, "                                                                 (наставник)", size=9, italic=True, after=8)
    para(doc, "Ознакомлен:", after=4)
    signatures(doc)
    para(doc, "                                                                 (работник)", size=9, italic=True, after=8)
    para(doc, "Согласовано руководителем структурного подразделения:", after=4)
    signatures(doc)
    para(doc, f"                                                                 {BOSS}", size=10)

    path = OUT / "03_Individualnyy_plan_nastavnichestva.docx"
    doc.save(path)
    return path


def build_readme() -> Path:
    text = f"""# Наставничество по Положению Белгидромета (PDF из кадров)

Сделано по скану **Положения о наставничестве** Белгидромета (приложения **1–3**), который выдали в отделе кадров.

## Файлы для печати

| Файл | Приложение | Назначение |
|------|------------|------------|
| `docx/01_Soglashenie_o_sotrudnichestve.docx` | 1 | Соглашение о сотрудничестве наставник ↔ работник |
| `docx/02_Programma_nastavnichestva.docx` | 2 | Программа наставничества подразделения |
| `docx/03_Individualnyy_plan_nastavnichestva.docx` | 3 | Индивидуальный план (+ оценка и характеристика) |

## Данные в бланках

- Организация: {ORG}
- Подразделение: {UNIT}
- Работник: {WORKER}, {WORKER_POS}
- Наставник: {MENTOR}, {MENTOR_POS}
- Руководитель: {BOSS}
- Период: {PERIOD}
- Основание: {BASIS}

В индивидуальном плане п. 1–5 и 13–14 — как в бланке кадров; п. 6–12 — конкретизация «освоения профессиональных навыков» (Python, Django, PostgreSQL, Oracle, нейросети/ML).

## Что сделать

1. Проверить ФИО, должности, даты.
2. Распечатать и подписать.
3. Отнести в кадры.

## Пересборка

```cmd
cd /d D:\\Work\\Nastavnichestvo\\po_polozheniyu_belgidromet
python build_forms.py
```
"""
    path = ROOT / "README.md"
    path.write_text(text, encoding="utf-8")
    return path


def main() -> None:
    for p in (build_agreement(), build_program(), build_plan(), build_readme()):
        print(p)


if __name__ == "__main__":
    main()

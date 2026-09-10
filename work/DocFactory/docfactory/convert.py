"""Офлайн-конвертеры: Markdown ↔ DOCX, DOCX ↔ PDF."""
from __future__ import annotations

import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from docfactory.engine import add_para, new_document, save_doc


class ConvertError(RuntimeError):
    pass


def _set_run_font(run, size: int = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def _add_inline(p, text: str, size: int = 12, bold: bool = False, italic: bool = False) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            _set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            _set_run_font(run, size=size)
            run.font.name = "Courier New"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = p.add_run(part[1:-1])
            _set_run_font(run, size=size, italic=True)
        else:
            run = p.add_run(part)
            _set_run_font(run, size=size, bold=bold, italic=italic)


def _parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.startswith("|"):
            break
        stripped = re.sub(r"\s", "", line)
        if re.match(r"^\|[\-:|]+\|$", stripped):
            i += 1
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
        i += 1
    return rows, i


def md_to_docx(src: Path, dst: Path | None = None) -> Path:
    """Markdown → DOCX (python-docx, офлайн)."""
    src = Path(src)
    if src.suffix.lower() != ".md":
        raise ConvertError("Ожидался файл .md")
    dst = Path(dst) if dst else src.with_suffix(".docx")
    lines = src.read_text(encoding="utf-8").splitlines()
    doc = new_document()
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        if not raw or raw.strip() == "---":
            i += 1
            continue
        if raw.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(12)
            _add_inline(p, raw[2:].strip(), size=16, bold=True)
            # force bold on whole title
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(16)
            i += 1
            continue
        if raw.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            _add_inline(p, raw[3:].strip(), size=14, bold=True)
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(14)
            i += 1
            continue
        if raw.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            _add_inline(p, raw[4:].strip(), size=13, bold=True)
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(13)
            i += 1
            continue
        if raw.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            _add_inline(p, raw[2:].strip(), size=11, italic=True)
            for run in p.runs:
                run.italic = True
            i += 1
            continue
        if re.match(r"^[-*]\s+", raw):
            content = re.sub(r"^[-*]\s+", "", raw)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            _add_inline(p, "• " + content, size=12)
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", raw)
        if m:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.75)
            p.paragraph_format.space_after = Pt(3)
            _add_inline(p, f"{m.group(1)}. {m.group(2)}", size=12)
            i += 1
            continue
        if raw.startswith("|"):
            rows, next_i = _parse_table(lines, i)
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=cols)
                table.style = "Table Grid"
                for r_idx, row in enumerate(rows):
                    for c_idx in range(cols):
                        cell_text = row[c_idx] if c_idx < len(row) else ""
                        cell = table.rows[r_idx].cells[c_idx]
                        cell.text = ""
                        clean = re.sub(r"\*\*([^*]+)\*\*", r"\1", cell_text)
                        clean = re.sub(r"`([^`]+)`", r"\1", clean)
                        run = cell.paragraphs[0].add_run(clean)
                        _set_run_font(run, size=10, bold=(r_idx == 0))
                doc.add_paragraph()
            i = next_i
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        _add_inline(p, raw, size=12)
        i += 1
    return save_doc(doc, dst)


def docx_to_md(src: Path, dst: Path | None = None) -> Path:
    """DOCX → Markdown (упрощённо: абзацы и таблицы)."""
    src = Path(src)
    if src.suffix.lower() != ".docx":
        raise ConvertError("Ожидался файл .docx")
    dst = Path(dst) if dst else src.with_suffix(".md")
    doc = Document(src)
    out: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            out.append("")
            continue
        style = (p.style.name or "") if p.style else ""
        if style.startswith("Heading 1"):
            out.append(f"# {text}")
        elif style.startswith("Heading 2"):
            out.append(f"## {text}")
        elif style.startswith("Heading 3"):
            out.append(f"### {text}")
        else:
            out.append(text)
    for table in doc.tables:
        out.append("")
        grid = [[c.text.strip().replace("\n", " ") for c in row.cells] for row in table.rows]
        if not grid:
            continue
        cols = max(len(r) for r in grid)
        for r in grid:
            while len(r) < cols:
                r.append("")
        out.append("| " + " | ".join(grid[0]) + " |")
        out.append("| " + " | ".join(["---"] * cols) + " |")
        for row in grid[1:]:
            out.append("| " + " | ".join(row) + " |")
        out.append("")
    dst.write_text("\n".join(out).strip() + "\n", encoding="utf-8")
    return dst


def _find_libreoffice() -> str | None:
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    # Common Windows installs
    candidates = [
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
    ]
    for c in candidates:
        if c.exists():
            return str(c)
    return None


def docx_to_pdf(src: Path, dst: Path | None = None) -> Path:
    """DOCX → PDF через LibreOffice (предпочтительно) или docx2pdf (Word на Windows)."""
    src = Path(src).resolve()
    if src.suffix.lower() != ".docx":
        raise ConvertError("Ожидался файл .docx")
    dst = Path(dst).resolve() if dst else src.with_suffix(".pdf")
    dst.parent.mkdir(parents=True, exist_ok=True)

    lo = _find_libreoffice()
    if lo:
        with tempfile.TemporaryDirectory() as tmp:
            cmd = [
                lo,
                "--headless",
                "--nologo",
                "--nofirststartwizard",
                "--convert-to",
                "pdf",
                "--outdir",
                tmp,
                str(src),
            ]
            proc = subprocess.run(cmd, capture_output=True, text=True)
            if proc.returncode != 0:
                raise ConvertError(
                    "LibreOffice не смог конвертировать в PDF.\n"
                    f"stdout: {proc.stdout}\nstderr: {proc.stderr}"
                )
            produced = Path(tmp) / (src.stem + ".pdf")
            if not produced.exists():
                pdfs = list(Path(tmp).glob("*.pdf"))
                if not pdfs:
                    raise ConvertError("LibreOffice отработал, но PDF не найден.")
                produced = pdfs[0]
            shutil.copy2(produced, dst)
        return dst

    # Fallback: docx2pdf (requires Microsoft Word on Windows)
    try:
        from docx2pdf import convert as docx2pdf_convert  # type: ignore
    except ImportError as exc:
        raise ConvertError(
            "Для DOCX→PDF нужен LibreOffice (soffice) или пакет docx2pdf + Microsoft Word.\n"
            "Установите LibreOffice: https://www.libreoffice.org/\n"
            "Либо: pip install docx2pdf  (только Windows с установленным Word)."
        ) from exc

    docx2pdf_convert(str(src), str(dst))
    if not dst.exists():
        raise ConvertError("docx2pdf завершился без создания PDF.")
    return dst


def _docx_text_stats(path: Path) -> tuple[int, int]:
    """Возвращает (число непробельных символов текста, число картинок)."""
    from zipfile import ZipFile
    import re

    path = Path(path)
    with ZipFile(path) as z:
        xml = z.read("word/document.xml")
        texts = re.findall(rb"<w:t[^>]*>([^<]*)</w:t>", xml)
        chars = sum(len(t.decode("utf-8", "ignore").strip()) for t in texts)
        media = [n for n in z.namelist() if n.startswith("word/media/")]
    return chars, len(media)


def _tesseract_available() -> bool:
    if shutil.which("tesseract"):
        return True
    # Common Windows install paths
    for p in (
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ):
        if p.exists():
            return True
    return False


def _tesseract_cmd() -> str | None:
    w = shutil.which("tesseract")
    if w:
        return w
    for p in (
        Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
        Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
    ):
        if p.exists():
            return str(p)
    return None


def pdf_ocr_pages(src: Path, lang: str = "rus+eng") -> list[str]:
    """OCR каждой страницы PDF → список текстов страниц."""
    src = Path(src)
    try:
        import pymupdf as fitz  # type: ignore
    except ImportError as exc:
        raise ConvertError("Для OCR нужен pymupdf (ставится с pdf2docx).") from exc
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore
        import io
    except ImportError as exc:
        raise ConvertError(
            "Для OCR установите: pip install pytesseract Pillow\n"
            "и программу Tesseract OCR: https://github.com/UB-Mannheim/tesseract/wiki"
        ) from exc

    cmd = _tesseract_cmd()
    if not cmd:
        raise ConvertError(
            "Tesseract OCR не найден.\n"
            "Windows: установите https://github.com/UB-Mannheim/tesseract/wiki\n"
            "при установке отметьте языки Russian + English.\n"
            "Затем перезапустите DocFactory."
        )
    pytesseract.pytesseract.tesseract_cmd = cmd

    doc = fitz.open(src)
    pages: list[str] = []
    try:
        for i in range(doc.page_count):
            page = doc[i]
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            img = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(img, lang=lang) or ""
            pages.append(text.strip())
    finally:
        doc.close()
    if not any(pages):
        raise ConvertError("OCR не распознал текст (пустой результат).")
    return pages


def pdf_ocr_to_docx(src: Path, dst: Path | None = None, lang: str = "rus+eng") -> Path:
    """PDF-скан → DOCX через OCR (Tesseract)."""
    src = Path(src)
    if src.suffix.lower() != ".pdf":
        raise ConvertError("Ожидался файл .pdf")
    dst = Path(dst) if dst else src.with_name(src.stem + "_ocr.docx")
    pages = pdf_ocr_pages(src, lang=lang)
    doc = new_document()
    for i, text in enumerate(pages):
        if i:
            doc.add_page_break()
        add_para(doc, f"— страница {i + 1} —", size=10, italic=True)
        if not text:
            add_para(doc, "[текст не распознан]")
            continue
        for line in text.splitlines():
            add_para(doc, line.rstrip() or " ")
    return save_doc(doc, dst)


def pdf_ocr_to_md(src: Path, dst: Path | None = None, lang: str = "rus+eng") -> Path:
    """PDF-скан → Markdown через OCR."""
    src = Path(src)
    dst = Path(dst) if dst else src.with_name(src.stem + "_ocr.md")
    pages = pdf_ocr_pages(src, lang=lang)
    parts: list[str] = [f"# OCR: {src.name}", ""]
    for i, text in enumerate(pages, start=1):
        parts.append(f"## Страница {i}")
        parts.append("")
        parts.append(text or "_[текст не распознан]_")
        parts.append("")
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text("\n".join(parts).strip() + "\n", encoding="utf-8")
    return dst


def pdf_to_docx(src: Path, dst: Path | None = None, *, force_ocr: bool = False) -> Path:
    """PDF → DOCX. Для сканов без текста — авто OCR (или force_ocr=True)."""
    src = Path(src)
    if src.suffix.lower() != ".pdf":
        raise ConvertError("Ожидался файл .pdf")
    dst = Path(dst) if dst else src.with_suffix(".docx")
    dst.parent.mkdir(parents=True, exist_ok=True)

    if force_ocr:
        return pdf_ocr_to_docx(src, dst)

    try:
        from pdf2docx import Converter  # type: ignore
    except ImportError as exc:
        # нет pdf2docx — сразу OCR
        return pdf_ocr_to_docx(src, dst)

    cv = Converter(str(src))
    try:
        cv.convert(str(dst))
    finally:
        cv.close()
    if not dst.exists():
        raise ConvertError("pdf2docx не создал файл DOCX.")

    chars, images = _docx_text_stats(dst)
    # Скан: почти нет текста, но есть картинки страниц
    if chars < 80 and images >= 1:
        ocr_dst = dst.with_name(dst.stem + "_ocr.docx")
        try:
            return pdf_ocr_to_docx(src, ocr_dst)
        except ConvertError:
            # оставляем «картиночный» docx, но предупреждаем
            raise ConvertError(
                "PDF похож на скан: в DOCX почти нет текста (только изображения страниц).\n"
                "Для распознавания установите Tesseract OCR "
                "(https://github.com/UB-Mannheim/tesseract/wiki) и пакеты:\n"
                "  pip install pytesseract Pillow\n"
                "Затем выберите режим «PDF → DOCX (OCR)»."
            )
    return dst


def convert_auto(src: Path, dst: Path | None = None, *, force_ocr: bool = False) -> Path:
    """Автовыбор направления по расширениям src/dst."""
    src = Path(src)
    if dst is None:
        mapping = {
            ".md": ".docx",
            ".docx": ".pdf",
            ".pdf": ".docx",
        }
        ext = mapping.get(src.suffix.lower())
        if not ext:
            raise ConvertError(f"Не знаю, во что конвертировать {src.suffix}")
        dst = src.with_suffix(ext)
    else:
        dst = Path(dst)

    pair = (src.suffix.lower(), dst.suffix.lower())
    if pair == (".md", ".docx"):
        return md_to_docx(src, dst)
    if pair == (".docx", ".md"):
        return docx_to_md(src, dst)
    if pair == (".docx", ".pdf"):
        return docx_to_pdf(src, dst)
    if pair == (".pdf", ".docx"):
        return pdf_to_docx(src, dst, force_ocr=force_ocr)
    if pair == (".pdf", ".md"):
        if force_ocr:
            return pdf_ocr_to_md(src, dst)
        tmp = dst.with_suffix(".tmp.docx")
        try:
            pdf_to_docx(src, tmp, force_ocr=False)
            # если получился OCR-файл с другим именем
            chars, _ = _docx_text_stats(tmp) if tmp.exists() else (0, 0)
            if chars < 80:
                return pdf_ocr_to_md(src, dst)
            return docx_to_md(tmp, dst)
        finally:
            if tmp.exists():
                tmp.unlink(missing_ok=True)
    if pair == (".md", ".pdf"):
        tmp = dst.with_suffix(".tmp.docx")
        md_to_docx(src, tmp)
        try:
            return docx_to_pdf(tmp, dst)
        finally:
            if tmp.exists():
                tmp.unlink(missing_ok=True)
    raise ConvertError(f"Неподдерживаемая пара: {pair[0]} → {pair[1]}")


def backend_status() -> dict[str, str]:
    """Статус доступных бэкендов конвертации."""
    status = {
        "md_docx": "ok (python-docx)",
        "docx_md": "ok (python-docx)",
        "docx_pdf": "нет LibreOffice / docx2pdf",
        "pdf_docx": "нет pdf2docx",
        "pdf_ocr": "нет Tesseract / pytesseract",
    }
    if _find_libreoffice():
        status["docx_pdf"] = f"ok (LibreOffice: {_find_libreoffice()})"
    else:
        try:
            import docx2pdf  # noqa: F401

            status["docx_pdf"] = "ok (docx2pdf / нужен Word)"
        except ImportError:
            pass
    try:
        import pdf2docx  # noqa: F401

        status["pdf_docx"] = "ok (pdf2docx)"
    except ImportError:
        pass
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401

        if _tesseract_available():
            status["pdf_ocr"] = f"ok (Tesseract: {_tesseract_cmd()})"
        else:
            status["pdf_ocr"] = "pytesseract есть, но Tesseract OCR не установлен"
    except ImportError:
        pass
    return status
"""Генераторы .docx по типам документов."""
from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docfactory.catalog import get_doc_type
from docfactory.engine import (
    add_bullets,
    add_grid_table,
    add_heading,
    add_kv_table,
    add_para,
    add_right_block,
    add_signature_block,
    add_title,
    new_document,
    save_doc,
    v,
)


def _lines(text: str) -> list[str]:
    return [ln.strip() for ln in (text or "").splitlines() if ln.strip()]


def _split_rows(text: str, cols: int) -> list[list[str]]:
    rows: list[list[str]] = []
    for ln in _lines(text):
        parts = [p.strip() for p in ln.split("|")] if "|" in ln else [ln]
        while len(parts) < cols:
            parts.append("")
        rows.append(parts[:cols])
    return rows or [[""] * cols]


def generate(doc_id: str, data: dict, out_dir: Path) -> Path:
    dtype = get_doc_type(doc_id)
    out_path = Path(out_dir) / dtype.filename
    generators = {
        "otchet_prodelannoy_raboty": _otchet_raboty,
        "otchet_po_proektu": _otchet_proekt,
        "otchet_za_period": _otchet_period,
        "otchet_promezhutochnyy": _otchet_promezh,
        "sluzhebnaya_zapiska": _sluzhebnaya,
        "dokladnaya_zapiska": _dokladnaya,
        "obyasnitelnaya": _obyasnitelnaya,
        "resume_cv": _resume,
        "kharakteristika": _kharakteristika,
        "akt_vypolnennyh_rabot": _akt_rabot,
        "akt_priema_peredachi": _akt_peredachi,
        "akt_proverki": _akt_proverki,
        "akt_spisaniya": _akt_spisaniya,
        "zayavlenie_pk": _zayavlenie_pk,
        "plan_pk": _plan_pk,
        "otchet_pk": _otchet_pk,
        "plan_nastavnichestva": _plan_nast,
        "zhurnal_nastavnichestva": _zhurnal_nast,
        "zayavlenie_otpusk": _zayavlenie,
        "protocol_soveshchaniya": _protokol,
    }
    doc = generators[doc_id](data)
    return save_doc(doc, out_path)


def _fill_paras(doc, text: str, placeholder: str) -> None:
    for ln in _lines(text) or [placeholder]:
        add_para(doc, ln)


def _otchet_raboty(data: dict):
    doc = new_document()
    add_title(doc, "ОТЧЁТ О ПРОДЕЛАННОЙ РАБОТЕ")
    add_kv_table(
        doc,
        [
            ("Организация", v(data, "org")),
            ("Подразделение", v(data, "podrazdelenie")),
            ("Исполнитель", f"{v(data, 'fio')}, {v(data, 'dolzhnost')}"),
            ("Отчётный период", v(data, "period")),
            ("Дата", v(data, "data")),
        ],
    )
    add_heading(doc, "1. Цель / задача периода")
    _fill_paras(doc, v(data, "cel", ""), "________________")
    add_heading(doc, "2. Выполненные работы")
    add_bullets(doc, _lines(v(data, "vypolneno", "")) or ["________________"])
    add_heading(doc, "3. Результаты")
    _fill_paras(doc, v(data, "rezultaty", ""), "________________")
    add_heading(doc, "4. Проблемы и риски")
    _fill_paras(doc, v(data, "problemy", ""), "________________")
    add_heading(doc, "5. Планы на следующий период")
    add_bullets(doc, _lines(v(data, "plany", "")) or ["________________"])
    add_signature_block(
        doc,
        [
            f"Исполнитель: _________________ / {v(data, 'fio')} /",
            f"Руководитель: _________________ / {v(data, 'rukovoditel')} /",
            v(data, "data"),
        ],
    )
    return doc


def _otchet_proekt(data: dict):
    doc = new_document()
    add_title(doc, "ОТЧЁТ О РЕЗУЛЬТАТАХ ПРОЕКТА")
    add_kv_table(
        doc,
        [
            ("Проект", v(data, "nazvanie")),
            ("Организация", v(data, "org")),
            ("Подразделение", v(data, "podrazdelenie")),
            ("Автор", f"{v(data, 'fio')}, {v(data, 'dolzhnost')}"),
            ("Сроки", v(data, "sroki")),
            ("Дата отчёта", v(data, "data")),
        ],
    )
    add_heading(doc, "1. Цель проекта")
    _fill_paras(doc, v(data, "cel", ""), "________________")
    add_heading(doc, "2. Задачи")
    add_bullets(doc, _lines(v(data, "zadachi", "")) or ["________________"])
    add_heading(doc, "3. Выполненные работы")
    add_bullets(doc, _lines(v(data, "vypolneno", "")) or ["________________"])
    add_heading(doc, "4. Достигнутые результаты")
    _fill_paras(doc, v(data, "rezultaty", ""), "________________")
    add_heading(doc, "5. Продукты / артефакты")
    add_bullets(doc, _lines(v(data, "produkty", "")) or ["________________"])
    add_heading(doc, "6. Выводы")
    _fill_paras(doc, v(data, "vyvody", ""), "________________")
    add_heading(doc, "7. Рекомендации")
    _fill_paras(doc, v(data, "rekomendacii", ""), "________________")
    add_signature_block(
        doc,
        [
            f"Автор: _________________ / {v(data, 'fio')} /",
            f"Согласовано: _________________ / {v(data, 'rukovoditel')} /",
            v(data, "data"),
        ],
    )
    return doc


def _otchet_period(data: dict):
    doc = new_document()
    add_title(doc, "ОТЧЁТ О РЕЗУЛЬТАТАХ РАБОТЫ ЗА ПЕРИОД")
    add_kv_table(
        doc,
        [
            ("Организация", v(data, "org")),
            ("Подразделение", v(data, "podrazdelenie")),
            ("Сотрудник", f"{v(data, 'fio')}, {v(data, 'dolzhnost')}"),
            ("Период", v(data, "period")),
            ("Дата", v(data, "data")),
        ],
    )
    add_heading(doc, "1. Ключевые показатели")
    add_grid_table(
        doc,
        ["Показатель", "План", "Факт"],
        _split_rows(v(data, "pokazateli", ""), 3),
    )
    add_heading(doc, "2. Основные работы")
    add_bullets(doc, _lines(v(data, "vypolneno", "")) or ["________________"])
    add_heading(doc, "3. Результаты")
    _fill_paras(doc, v(data, "rezultaty", ""), "________________")
    add_heading(doc, "4. Выводы")
    _fill_paras(doc, v(data, "vyvody", ""), "________________")
    add_signature_block(
        doc,
        [
            f"_________________ / {v(data, 'fio')} /",
            f"Руководитель: _________________ / {v(data, 'rukovoditel')} /",
            v(data, "data"),
        ],
    )
    return doc


def _otchet_promezh(data: dict):
    doc = new_document()
    add_title(doc, "ПРОМЕЖУТОЧНЫЙ ОТЧЁТ ПО ПРОЕКТУ")
    add_kv_table(
        doc,
        [
            ("Проект", v(data, "nazvanie")),
            ("Организация", v(data, "org")),
            ("Автор", f"{v(data, 'fio')}, {v(data, 'dolzhnost')}"),
            ("Дата статуса", v(data, "data_statusa")),
            ("Прогресс", v(data, "progress")),
        ],
    )
    add_heading(doc, "1. Сделано с прошлого отчёта")
    add_bullets(doc, _lines(v(data, "sdelano", "")) or ["________________"])
    add_heading(doc, "2. В работе сейчас")
    add_bullets(doc, _lines(v(data, "v_rabote", "")) or ["________________"])
    add_heading(doc, "3. Блокеры / риски")
    _fill_paras(doc, v(data, "blockers", ""), "________________")
    add_heading(doc, "4. Ближайшие шаги")
    add_bullets(doc, _lines(v(data, "next_steps", "")) or ["________________"])
    add_heading(doc, "5. Нужна помощь / решения")
    _fill_paras(doc, v(data, "nuzhna_pomoshch", ""), "________________")
    add_signature_block(
        doc,
        [
            f"Автор: _________________ / {v(data, 'fio')} /",
            f"Руководитель: _________________ / {v(data, 'rukovoditel')} /",
            v(data, "data_statusa"),
        ],
    )
    return doc


def _sluzhebnaya(data: dict):
    doc = new_document()
    add_right_block(
        doc,
        [v(data, "komu"), f"от {v(data, 'ot_kogo')}", v(data, "org"), v(data, "podrazdelenie")],
    )
    add_title(doc, "СЛУЖЕБНАЯ ЗАПИСКА")
    add_para(doc, v(data, "data"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_para(doc, f"Тема: {v(data, 'tema')}", bold=True)
    _fill_paras(doc, v(data, "tekst", ""), "[текст записки]")
    add_heading(doc, "Прошу / предлагаю")
    _fill_paras(doc, v(data, "prosba", ""), "________________")
    add_signature_block(doc, [f"_________________ / {v(data, 'ot_kogo')} /", v(data, "data")])
    return doc


def _dokladnaya(data: dict):
    doc = new_document()
    add_right_block(doc, [v(data, "komu"), f"от {v(data, 'ot_kogo')}", v(data, "org")])
    add_title(doc, "ДОКЛАДНАЯ ЗАПИСКА")
    add_para(doc, v(data, "data"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_para(doc, f"Тема: {v(data, 'tema')}", bold=True)
    _fill_paras(doc, v(data, "tekst", ""), "[содержание]")
    add_heading(doc, "Выводы и предложения")
    _fill_paras(doc, v(data, "vyvody", ""), "________________")
    add_signature_block(doc, [f"_________________ / {v(data, 'ot_kogo')} /", v(data, "data")])
    return doc


def _obyasnitelnaya(data: dict):
    doc = new_document()
    add_right_block(doc, [v(data, "komu"), f"от {v(data, 'fio')}, {v(data, 'dolzhnost')}", v(data, "org")])
    add_title(doc, "ОБЪЯСНИТЕЛЬНАЯ ЗАПИСКА")
    add_para(doc, v(data, "data"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_para(doc, f"По поводу: {v(data, 'po_povodu')}", bold=True)
    _fill_paras(doc, v(data, "tekst", ""), "[объяснение]")
    add_signature_block(doc, [f"_________________ / {v(data, 'fio')} /", v(data, "data")])
    return doc


def _resume(data: dict):
    doc = new_document()
    add_title(doc, v(data, "fio", "ФАМИЛИЯ ИМЯ ОТЧЕСТВО"))
    add_para(
        doc,
        f"{v(data, 'dolzhnost_cel')}  |  {v(data, 'gorod')}  |  {v(data, 'telefon')}  |  {v(data, 'email')}",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=11,
    )
    add_heading(doc, "О себе")
    _fill_paras(doc, v(data, "o_sebe", ""), "________________")
    add_heading(doc, "Опыт работы")
    for ln in _lines(v(data, "opyt", "")) or ["________________"]:
        add_para(doc, f"• {ln}", left_indent_cm=0.3)
    add_heading(doc, "Образование")
    for ln in _lines(v(data, "obrazovanie", "")) or ["________________"]:
        add_para(doc, f"• {ln}", left_indent_cm=0.3)
    add_heading(doc, "Навыки")
    skills = v(data, "navyki", "")
    if "," in skills and "\n" not in skills.strip():
        add_bullets(doc, [s.strip() for s in skills.split(",") if s.strip()])
    else:
        add_bullets(doc, _lines(skills) or ["________________"])
    add_heading(doc, "Языки")
    add_para(doc, v(data, "yazyki", "________________"))
    return doc


def _kharakteristika(data: dict):
    doc = new_document()
    add_title(doc, "ХАРАКТЕРИСТИКА")
    add_para(
        doc,
        f"на {v(data, 'fio')}, {v(data, 'dolzhnost')}, {v(data, 'org')}, {v(data, 'podrazdelenie')}.",
    )
    add_para(doc, f"Период работы: {v(data, 'period')}.")
    _fill_paras(doc, v(data, "tekst", ""), "[текст характеристики]")
    add_heading(doc, "Рекомендация")
    _fill_paras(doc, v(data, "rekomendaciya", ""), "________________")
    add_signature_block(
        doc,
        [f"_________________ / {v(data, 'rukovoditel')} /", v(data, "data"), "М.П."],
    )
    return doc


def _akt_rabot(data: dict):
    doc = new_document()
    add_title(doc, f"АКТ ВЫПОЛНЕННЫХ РАБОТ № {v(data, 'nomer')}")
    add_para(doc, v(data, "data"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_kv_table(
        doc,
        [
            ("Организация", v(data, "org")),
            ("Заказчик", v(data, "zakazchik")),
            ("Исполнитель", v(data, "ispolnitel")),
            ("Период", v(data, "period")),
        ],
    )
    add_heading(doc, "Перечень работ")
    add_grid_table(doc, ["Наименование", "Ед.", "Кол-во", "Сумма"], _split_rows(v(data, "raboty", ""), 4))
    add_para(doc, f"Итого: {v(data, 'itogo')}", bold=True)
    for ln in _lines(v(data, "primechanie", "")):
        add_para(doc, ln, italic=True)
    add_para(doc, "Работы выполнены в полном объёме, стороны претензий не имеют.")
    add_signature_block(
        doc,
        [
            f"Заказчик: _________________ / {v(data, 'zakazchik')} /",
            f"Исполнитель: _________________ / {v(data, 'ispolnitel')} /",
        ],
    )
    return doc


def _akt_peredachi(data: dict):
    doc = new_document()
    add_title(doc, f"АКТ ПРИЁМА-ПЕРЕДАЧИ № {v(data, 'nomer')}")
    add_para(doc, v(data, "data"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_kv_table(
        doc,
        [
            ("Организация", v(data, "org")),
            ("Основание", v(data, "osnovanie")),
            ("Сдал", v(data, "sdal")),
            ("Принял", v(data, "prinyal")),
        ],
    )
    add_heading(doc, "Перечень")
    items = _lines(v(data, "perechen", "")) or ["________________"]
    add_grid_table(doc, ["№", "Наименование / описание"], [[str(i + 1), ln] for i, ln in enumerate(items)])
    add_signature_block(
        doc,
        [
            f"Сдал: _________________ / {v(data, 'sdal')} /",
            f"Принял: _________________ / {v(data, 'prinyal')} /",
        ],
    )
    return doc


def _akt_proverki(data: dict):
    doc = new_document()
    add_title(doc, f"АКТ ПРОВЕРКИ № {v(data, 'nomer')}")
    add_para(doc, v(data, "data"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_kv_table(
        doc,
        [
            ("Организация", v(data, "org")),
            ("Подразделение", v(data, "podrazdelenie")),
            ("Объект", v(data, "obiekt")),
            ("Цель", v(data, "cel")),
        ],
    )
    add_heading(doc, "Состав комиссии")
    add_bullets(doc, _lines(v(data, "komissiya", "")) or ["________________"])
    add_heading(doc, "Выявлено")
    _fill_paras(doc, v(data, "vyyavleno", ""), "________________")
    add_heading(doc, "Заключение")
    _fill_paras(doc, v(data, "zaklyuchenie", ""), "________________")
    add_signature_block(
        doc,
        [
            "Члены комиссии:",
            "_________________ / _________________ /",
            "_________________ / _________________ /",
        ],
    )
    return doc


def _akt_spisaniya(data: dict):
    doc = new_document()
    add_title(doc, f"АКТ СПИСАНИЯ № {v(data, 'nomer')}")
    add_para(doc, v(data, "data"), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_kv_table(doc, [("Организация", v(data, "org")), ("Основание", v(data, "osnovanie"))])
    add_heading(doc, "Комиссия")
    add_bullets(doc, _lines(v(data, "komissiya", "")) or ["________________"])
    add_heading(doc, "К списанию")
    add_grid_table(doc, ["Наименование", "Кол-во", "Причина"], _split_rows(v(data, "perechen", ""), 3))
    add_heading(doc, "Заключение")
    _fill_paras(doc, v(data, "zaklyuchenie", ""), "________________")
    add_signature_block(
        doc,
        [
            "Председатель комиссии: _________________ / _________________ /",
            "Члены комиссии: _________________ / _________________ /",
        ],
    )
    return doc


def _zayavlenie_pk(data: dict):
    doc = new_document()
    add_right_block(doc, [v(data, "komu"), f"от {v(data, 'fio')}, {v(data, 'dolzhnost')}", v(data, "org")])
    add_title(doc, "ЗАЯВЛЕНИЕ")
    add_para(
        doc,
        f"Прошу направить меня на повышение квалификации по программе «{v(data, 'programma')}» "
        f"(организатор: {v(data, 'organizator')}, сроки: {v(data, 'sroki')}).",
    )
    add_heading(doc, "Обоснование")
    _fill_paras(doc, v(data, "obosnovanie", ""), "________________")
    add_signature_block(doc, [f"_________________ / {v(data, 'fio')} /", v(data, "data")])
    return doc


def _plan_pk(data: dict):
    doc = new_document()
    add_title(doc, "ПЛАН ПОВЫШЕНИЯ КВАЛИФИКАЦИИ")
    add_kv_table(
        doc,
        [
            ("Организация", v(data, "org")),
            ("Работник", f"{v(data, 'fio')}, {v(data, 'dolzhnost')}"),
            ("Период", v(data, "god")),
        ],
    )
    add_heading(doc, "Цель")
    _fill_paras(doc, v(data, "cel", ""), "________________")
    add_heading(doc, "Мероприятия")
    add_grid_table(
        doc,
        ["Дата / срок", "Тема", "Форма", "Ожидаемый результат"],
        _split_rows(v(data, "meropriyatiya", ""), 4),
    )
    add_signature_block(
        doc,
        [
            f"Работник: _________________ / {v(data, 'fio')} /",
            f"Руководитель: _________________ / {v(data, 'rukovoditel')} /",
        ],
    )
    return doc


def _otchet_pk(data: dict):
    doc = new_document()
    add_title(doc, "ОТЧЁТ О ПОВЫШЕНИИ КВАЛИФИКАЦИИ")
    add_kv_table(
        doc,
        [
            ("ФИО", v(data, "fio")),
            ("Должность", v(data, "dolzhnost")),
            ("Организация", v(data, "org")),
            ("Программа", v(data, "programma")),
            ("Сроки", v(data, "sroki")),
            ("Место", v(data, "gde")),
        ],
    )
    add_heading(doc, "Содержание обучения")
    _fill_paras(doc, v(data, "soderzhanie", ""), "________________")
    add_heading(doc, "Применение на рабочем месте")
    _fill_paras(doc, v(data, "primenenie", ""), "________________")
    add_signature_block(doc, [f"_________________ / {v(data, 'fio')} /", v(data, "data")])
    return doc


def _plan_nast(data: dict):
    doc = new_document()
    add_title(doc, "ПЛАН ИНДИВИДУАЛЬНОГО НАСТАВНИЧЕСТВА")
    add_para(doc, f"Основание: приказ {v(data, 'prikaz')} «О наставничестве».", italic=True)
    add_para(doc, "УТВЕРЖДАЮ", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_para(doc, f"_________________ / {v(data, 'rukovoditel')} /", align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_kv_table(
        doc,
        [
            ("Организация", v(data, "org")),
            ("Подразделение", v(data, "podrazdelenie")),
            ("Наставляемый", v(data, "nastavlyaemyj")),
            ("Наставник", v(data, "nastavnik")),
            ("Период", v(data, "period")),
        ],
    )
    add_heading(doc, "Цель")
    _fill_paras(doc, v(data, "cel", ""), "________________")
    add_heading(doc, "Темы подготовки")
    add_para(doc, v(data, "temy"))
    add_heading(doc, "Этапы")
    add_grid_table(
        doc,
        ["Срок", "Содержание", "Ожидаемый результат"],
        _split_rows(v(data, "etapy", ""), 3),
    )
    add_signature_block(
        doc,
        [
            f"Наставляемый: _________________ / {v(data, 'nastavlyaemyj')} /",
            f"Наставник: _________________ / {v(data, 'nastavnik')} /",
        ],
    )
    return doc


def _zhurnal_nast(data: dict):
    doc = new_document()
    add_title(doc, "ЖУРНАЛ НАСТАВНИЧЕСТВА")
    add_para(doc, f"Основание: приказ {v(data, 'prikaz')}.", italic=True)
    add_kv_table(
        doc,
        [
            ("Наставляемый", v(data, "nastavlyaemyj")),
            ("Наставник", v(data, "nastavnik")),
            ("Период", v(data, "period")),
            ("Организация", v(data, "org")),
        ],
    )
    add_heading(doc, "Записи")
    add_grid_table(
        doc,
        ["Дата", "Тема", "Содержание", "Результат", "Задание"],
        _split_rows(v(data, "zapisi", ""), 5),
        size=9,
    )
    add_signature_block(
        doc,
        [
            f"Наставник: _________________ / {v(data, 'nastavnik')} /",
            f"Наставляемый: _________________ / {v(data, 'nastavlyaemyj')} /",
        ],
    )
    return doc


def _zayavlenie(data: dict):
    doc = new_document()
    add_right_block(doc, [v(data, "komu"), f"от {v(data, 'fio')}, {v(data, 'dolzhnost')}", v(data, "org")])
    add_title(doc, "ЗАЯВЛЕНИЕ")
    _fill_paras(doc, v(data, "tekst", ""), "Прошу предоставить...")
    add_signature_block(doc, [f"_________________ / {v(data, 'fio')} /", v(data, "data")])
    return doc


def _protokol(data: dict):
    doc = new_document()
    add_title(doc, f"ПРОТОКОЛ СОВЕЩАНИЯ № {v(data, 'nomer')}")
    add_para(doc, f"{v(data, 'org')}, {v(data, 'podrazdelenie')}")
    add_para(doc, f"Дата: {v(data, 'data')}")
    add_para(doc, f"Тема: {v(data, 'tema')}", bold=True)
    add_heading(doc, "Участники")
    add_bullets(doc, _lines(v(data, "uchastniki", "")) or ["________________"])
    add_heading(doc, "Повестка")
    add_bullets(doc, _lines(v(data, "povestka", "")) or ["________________"])
    add_heading(doc, "Решения / поручения")
    add_bullets(doc, _lines(v(data, "resheniya", "")) or ["________________"])
    add_signature_block(
        doc,
        [
            f"Председатель: _________________ / {v(data, 'predsedatel')} /",
            f"Секретарь: _________________ / {v(data, 'sekretar')} /",
        ],
    )
    return doc

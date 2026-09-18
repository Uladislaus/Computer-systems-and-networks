"""Общие хелперы для генерации .docx (Times New Roman, таблицы, подписи)."""
from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def new_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc


def _font(run, size: int = 12, bold: bool = False, italic: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(
    doc: Document,
    text: str,
    *,
    size: int = 12,
    bold: bool = False,
    italic: bool = False,
    align=None,
    space_after: int = 6,
    left_indent_cm: float | None = None,
) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if left_indent_cm is not None:
        p.paragraph_format.left_indent = Cm(left_indent_cm)
    run = p.add_run(text)
    _font(run, size=size, bold=bold, italic=italic)


def add_title(doc: Document, text: str) -> None:
    add_para(doc, text, size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)


def add_heading(doc: Document, text: str, level: int = 2) -> None:
    size = 14 if level == 2 else 13
    add_para(doc, text, size=size, bold=True, space_after=8)


def add_right_block(doc: Document, lines: Sequence[str]) -> None:
    for line in lines:
        add_para(doc, line, size=12, align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    add_para(doc, "", space_after=8)


def add_kv_table(doc: Document, rows: Sequence[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.text = ""
        c1.text = ""
        r0 = c0.paragraphs[0].add_run(k)
        _font(r0, size=11, bold=True)
        r1 = c1.paragraphs[0].add_run(v)
        _font(r1, size=11)
    doc.add_paragraph()


def add_grid_table(doc: Document, header: Sequence[str], rows: Sequence[Sequence[str]], size: int = 10) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        _font(run, size=size, bold=True)
    for i, row in enumerate(rows, start=1):
        for j in range(len(header)):
            cell = table.rows[i].cells[j]
            cell.text = ""
            val = row[j] if j < len(row) else ""
            run = cell.paragraphs[0].add_run(val)
            _font(run, size=size)
    doc.add_paragraph()


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        add_para(doc, f"• {item}", size=12, space_after=3, left_indent_cm=0.5)


def add_signature_block(doc: Document, lines: Sequence[str]) -> None:
    add_para(doc, "", space_after=6)
    for line in lines:
        add_para(doc, line, size=12, space_after=8)


def v(data: dict, key: str, default: str = "________________") -> str:
    val = (data.get(key) or "").strip()
    return val if val else default


def save_doc(doc: Document, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
